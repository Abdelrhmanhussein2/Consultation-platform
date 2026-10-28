import io
import re
import uuid
try:
    import docx
except ImportError:
    docx = None
from helpers.neo4j_db import neo4j_db

class DocxParserService:
    @staticmethod
    def extract_docx_text(file_bytes: bytes) -> str:
        """
        Extracts raw text from both paragraphs and tables of a DOCX file.
        """
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        
        # 1. Extract from paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                full_text.append(para.text)
                
        # 2. Extract from tables (useful if laws/articles are formatted in tables)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        full_text.append(cell.text)
                        
        return '\n'.join(full_text)

    ARABIC_LETTERS = {
        'أ': 1, 'ا': 1, 'إ': 1, 'آ': 1,
        'ب': 2, 'ج': 3, 'د': 4, 'هـ': 5, 'ه': 5,
        'و': 6, 'ز': 7, 'ح': 8, 'ط': 9, 'ي': 10, 'ى': 10,
        'ك': 11, 'ل': 12, 'م': 13, 'ن': 14, 'س': 15, 'ع': 16,
        'ف': 17, 'ص': 18, 'ق': 19, 'ر': 20, 'ش': 21, 'ت': 22,
        'ث': 23, 'خ': 24, 'ذ': 25, 'ض': 26, 'ظ': 27, 'غ': 28
    }

    ARABIC_ORDINALS = {
        'أولاً': 1, 'اولا': 1, 'أولا': 1,
        'ثانياً': 2, 'ثانيا': 2,
        'ثالثاً': 3, 'ثالثا': 3,
        'رابعاً': 4, 'رابعا': 4,
        'خامساً': 5, 'خامسا': 5,
        'سادساً': 6, 'سادسا': 6,
        'سابعاً': 7, 'سابعا': 7,
        'ثامناً': 8, 'ثامنا': 8,
        'تاسعاً': 9, 'تاسعا': 9,
        'عاشراً': 10, 'عاشرا': 10
    }

    @classmethod
    def parse_hierarchical_prefix(cls, text: str):
        if not text:
            return None
        t = text.strip()
        ord_match = re.match(r'^(أولاً|اولا|أولا|ثانياً|ثانيا|ثالثاً|ثالثا|رابعاً|رابعا|خامساً|خامسا|سادساً|سادسا|سابعاً|سابعا|ثامناً|ثامنا|تاسعاً|تاسعا|عاشراً|عاشرا)\s*[:\-–—\.]', t)
        if ord_match:
            return {'type': 'ordinal', 'val': cls.ARABIC_ORDINALS.get(ord_match.group(1), 1), 'raw': ord_match.group(0)}
        
        letter_match = re.match(r'^(?:([أ-ي]|هـ)\s*[-–—\.\:]|\(([أ-ي]|هـ)\)|([أ-ي]|هـ)\))', t)
        if letter_match:
            l = letter_match.group(1) or letter_match.group(2) or letter_match.group(3)
            return {'type': 'letter', 'val': cls.ARABIC_LETTERS.get(l, 0), 'raw': letter_match.group(0)}
            
        num_match = re.match(r'^(?:(\d+|[٠-٩]+)\s*[-–—\.\:]|\((\d+|[٠-٩]+)\)|(\d+|[٠-٩]+)\))', t)
        if num_match:
            n = num_match.group(1) or num_match.group(2) or num_match.group(3)
            trans = str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789')
            return {'type': 'number', 'val': int(n.translate(trans)), 'raw': num_match.group(0)}
            
        return None

    @classmethod
    def assign_hierarchical_levels(cls, lines: list) -> list:
        """
        Assigns progressive indentation levels (0, 1, 2, 3...) to Arabic legal clauses.
        Returns list of dicts: [{'text': line, 'level': lvl, 'prefix': prefix}]
        """
        stack = []
        results = []
        for line in lines:
            line_str = line if isinstance(line, str) else str(line)
            prefix = cls.parse_hierarchical_prefix(line_str)
            if not prefix:
                level = stack[-1]['level'] if stack else 0
                results.append({'text': line_str, 'level': level, 'prefix': None})
                continue
                
            if not stack:
                level = 0
                stack.append({'type': prefix['type'], 'val': prefix['val'], 'level': 0})
            else:
                top = stack[-1]
                if top['type'] == prefix['type'] and prefix['val'] > top['val']:
                    level = top['level']
                    top['val'] = prefix['val']
                elif prefix['val'] == 1 and top['type'] != prefix['type']:
                    level = top['level'] + 1
                    stack.append({'type': prefix['type'], 'val': prefix['val'], 'level': level})
                else:
                    match_idx = -1
                    for i in range(len(stack) - 1, -1, -1):
                        if stack[i]['type'] == prefix['type'] and prefix['val'] > stack[i]['val']:
                            match_idx = i
                            break
                    if match_idx != -1:
                        del stack[match_idx + 1:]
                        level = stack[match_idx]['level']
                        stack[match_idx]['val'] = prefix['val']
                    else:
                        root_idx = next((i for i, s in enumerate(stack) if s['type'] == prefix['type']), -1)
                        if root_idx != -1:
                            del stack[root_idx + 1:]
                            level = stack[root_idx]['level']
                            stack[root_idx]['val'] = prefix['val']
                        else:
                            stack.clear()
                            level = 0
                            stack.append({'type': prefix['type'], 'val': prefix['val'], 'level': 0})
            results.append({'text': line_str, 'level': level, 'prefix': prefix})
        return results

    @classmethod
    def format_hierarchical_text(cls, text: str, indent_str: str = "    ") -> str:
        """
        Formats legal text lines with hierarchical progressive indentation.
        """
        if not text:
            return ""
        lines = [l.strip() for l in text.split('\n') if l.strip()]
        hierarchical = cls.assign_hierarchical_levels(lines)
        formatted_lines = []
        for h in hierarchical:
            indent = indent_str * h['level']
            formatted_lines.append(f"{indent}{h['text']}")
        return '\n'.join(formatted_lines)

    @classmethod
    def resolve_law_name_to_id(cls, law_name: str) -> str:
        """
        Resolves an Arabic law name to its corresponding law_id by searching in Neo4j.
        """
        name_clean = law_name.lower().strip()
        name_clean = re.sub(r'^(قانون|ال)|\s+', '', name_clean)
        
        # Static lookup mappings for common laws
        if "ضريبةالدخل" in name_clean or "ضريبهالدخل" in name_clean:
            return "law_34_2014"
        elif "جمارك" in name_clean:
            return "law_20_1998"
            
        # Dynamic search in Neo4j
        query = """
        MATCH (l:Law)
        WHERE l.title CONTAINS $keyword OR $keyword CONTAINS l.title
        RETURN l.law_id AS law_id
        LIMIT 1
        """
        try:
            keyword = law_name.replace("قانون", "").strip()
            if not keyword:
                return None
            with neo4j_db.get_session() as session:
                res = session.execute_read(lambda tx: tx.run(query, keyword=keyword).single())
                if res:
                    return res["law_id"]
        except Exception:
            pass
        return None

    @classmethod
    def parse_law_docx(cls, file_bytes: bytes, law_id: str, version_name: str, effective_from: str) -> dict:
        """
        Parses a law DOCX file and extracts ArticleVersions, Paragraphs, and Items.
        - Articles are matched using patterns like: المادة 1-, المادة (5), المادة 12:
        - Paragraphs are matched using Arabic alphabetical letters: أ-, ب-, ج)
        - Items are matched using digit numbers: 1-, 2-, 3)
        """
        text = cls.extract_docx_text(file_bytes)
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        
        # Regex patterns for Arabic legal structure
        article_pattern = re.compile(r'^المادة\s*\(?(\d+)\)?[\-\:\s\.]*')
        paragraph_pattern = re.compile(r'^([أ-ي])[\-\:\s\.\)]+\s*(.*)')
        item_pattern = re.compile(r'^(\d+)[\-\:\s\.\)]+\s*(.*)')
        
        articles = []
        paragraphs = []
        items = []
        
        current_article = None
        current_paragraph = None
        
        for line in lines:
            # 1. Match Article
            art_match = article_pattern.match(line)
            if art_match:
                art_num = int(art_match.group(1))
                art_version_id = f"{law_id}_art_{art_num}_v_{version_name}"
                
                # Check for explicit title on the same line after "المادة 1: ..."
                remainder = line[art_match.end():].strip()
                remainder_clean = re.sub(r'^[\-\:\.\s]+', '', remainder).strip()
                
                current_article = {
                    "version_id": art_version_id,
                    "text": line,
                    "title": remainder_clean if len(remainder_clean) > 2 else f"المادة {art_num}",
                    "effective_from": effective_from,
                    "status": "active",
                    "number": art_num,
                    "law_version_id": f"{law_id}_v_{version_name}"
                }
                articles.append(current_article)
                current_paragraph = None
                continue
                
            # 2. Match Paragraph (only if inside an active article)
            if current_article:
                p_match = paragraph_pattern.match(line)
                if p_match:
                    p_letter = p_match.group(1)
                    p_id = f"{current_article['version_id']}_p_{p_letter}"
                    
                    current_paragraph = {
                        "paragraph_id": p_id,
                        "letter": p_letter,
                        "text": line,
                        "article_version_id": current_article["version_id"]
                    }
                    paragraphs.append(current_paragraph)
                    continue
                    
                # 3. Match Item (only if inside an active paragraph)
                if current_paragraph:
                    i_match = item_pattern.match(line)
                    if i_match:
                        i_num = int(i_match.group(1))
                        i_id = f"{current_paragraph['paragraph_id']}_i_{i_num}"
                        
                        item_node = {
                            "item_id": i_id,
                            "number": i_num,
                            "text": line,
                            "paragraph_id": current_paragraph["paragraph_id"]
                        }
                        items.append(item_node)
                        continue
                        
                # 4. Multiline logic: if no structure matches, append text to the current open node
                if current_paragraph:
                    current_paragraph["text"] += "\n" + line
                elif current_article:
                    # Guard: only append if we are inside an article (skip preamble text)
                    current_article["text"] += "\n" + line
                # else: preamble text before first article — safely ignored

        # Enrich articles with hierarchical clauses and formatted indentation
        for art in articles:
            lines_in_art = [l.strip() for l in art["text"].split('\n') if l.strip()]
            hierarchical = cls.assign_hierarchical_levels(lines_in_art)
            art["hierarchical_clauses"] = hierarchical
            art["formatted_text"] = cls.format_hierarchical_text(art["text"])

        return {
            "articles": articles,
            "paragraphs": paragraphs,
            "items": items
        }

    @classmethod
    def parse_judgment_docx(
        cls,
        file_bytes: bytes,
        metadata: dict = None,
        default_law_id: str = None,
        default_version_name: str = None
    ) -> dict:
        """
        Parses a judgment DOCX, extracts court metadata using Regex, and finds legal citations.
        """
        text = cls.extract_docx_text(file_bytes)
        
        extracted = {
            "ruling_id": None,
            "case_number": None,
            "ruling_number": None,
            "ruling_year": None,
            "court": None,
            "court_type": "cassation",
            "date": None,
            "outcome": None,
            "subject": None,
            "title": None,
            "full_text": text
        }
        
        # 1. Extract metadata via Regex
        case_match = re.search(r'(?:قضية رقم|رقم القضية|الدعوى رقم|رقم الدعوى)\s*[:\-]?\s*([\d\s]+/[\s\d]+)', text)
        if case_match:
            extracted["case_number"] = case_match.group(1).replace(" ", "")
            
        ruling_match = re.search(r'(?:حكم رقم|قرار رقم|الحكم رقم|القرار رقم|رقم الحكم)\s*[:\-]?\s*(\d+)(?:\s*/\s*|\s+لسنة\s+)(\d+)', text)
        if not ruling_match:
            ruling_match = re.search(r'رقم\s*(\d+)\s+لسنة\s+(\d+)', text)
        if ruling_match:
            extracted["ruling_number"] = int(ruling_match.group(1))
            extracted["ruling_year"] = int(ruling_match.group(2))
            
        court_std_match = re.search(r'(المحكمة\s+الإدارية\s+العليا|محكمة\s+التمييز\s+الأردنية|محكمة\s+التمييز|محكمة\s+البداية\s+الضريبية|محكمة\s+الاستئناف\s+الضريبية)', text)
        if court_std_match:
            extracted["court"] = court_std_match.group(1).strip()
        else:
            court_match = re.search(r'محكمة\s+([^\s،؛\.\n\r][^،؛\.\n\r]+)', text)
            if court_match:
                extracted["court"] = "محكمة " + court_match.group(1).strip()
            
        date_match = re.search(r'(?:تاريخ الفصل|تاريخ النطق|تاريخ الصدور|تاريخ)\s*[:\-]?\s*([\d\-\/\.\s]{8,12})', text)
        if not date_match:
            date_match = re.search(r'بتاريخ\s*(\d{1,2}\s*[\/\-]\s*\d{1,2}\s*[\/\-]\s*\d{2,4})', text)
        if date_match:
            extracted["date"] = date_match.group(1).replace(" ", "").strip()
            
        subject_match = re.search(r'(?:الموضوع|موضوع الدعوى|موضوع الدعــوى)\s*[:\-]?\s*([^\n\r]+)', text)
        if subject_match:
            extracted["subject"] = subject_match.group(1).strip()
        else:
            filing_match = re.search(r'(بتاريخ\s+[\d\s\/]+(?:تقدم|قدم|أقامت|أقام)[^\n\r]+)', text)
            if filing_match:
                extracted["subject"] = filing_match.group(1).strip()
            
        outcome_match = re.search(r'(?:القرار|النتيجة)\s*[:\-]\s*([^\n\r]+)', text)
        if outcome_match:
            extracted["outcome"] = outcome_match.group(1).strip()

        # Override with user metadata if provided
        if metadata:
            for k, v in metadata.items():
                if v is not None:
                    extracted[k] = v

        # Construct ruling_id if not present
        if not extracted["ruling_id"]:
            if extracted["ruling_number"] and extracted["ruling_year"]:
                extracted["ruling_id"] = f"ruling_{extracted['ruling_number']}_{extracted['ruling_year']}"
            else:
                extracted["ruling_id"] = f"ruling_{uuid.uuid4().hex[:8]}"

        # Deduce court_type
        court_name = (extracted["court"] or "").lower()
        if "تمييز" in court_name:
            extracted["court_type"] = "cassation"
        elif "استئناف" in court_name and "ضريب" in court_name:
            extracted["court_type"] = "appeal_tax"
        elif "بداية" in court_name and "ضريب" in court_name:
            extracted["court_type"] = "tax_first"
            
        if not extracted["title"]:
            extracted["title"] = f"قرار {extracted['court'] or 'المحكمة'} رقم {extracted['ruling_number'] or ''} لسنة {extracted['ruling_year'] or ''}"

        # 2. Extract Citations
        # Pattern matches "المادة 11/و/2 من قانون..." or "المادة (70) من القانون"
        citation_pattern = re.compile(
            r'المادة\s*\(?(\d+)\)?(?:\s*/\s*\(?([أ-ي])\)?)*(?:\s*/\s*\(?(\d+)\)?)*(?:\s+من\s+(قانون\s+[^\s،؛\.][^،؛\.]+))?',
            re.IGNORECASE
        )
        
        citations = []
        for match in citation_pattern.finditer(text):
            art_num = match.group(1)
            p_letter = match.group(2)
            item_num = match.group(3)
            law_name = match.group(4)
            
            # Grab context around citation for Neo4j property storage
            start = max(0, match.start() - 30)
            end = min(len(text), match.end() + 30)
            citation_context = text[start:end].strip().replace("\n", " ")
            
            resolved_law_id = default_law_id
            if law_name:
                resolved_law_id = cls.resolve_law_name_to_id(law_name) or default_law_id
                
            if not resolved_law_id:
                continue
                
            resolved_version_name = default_version_name or "amended"
            
            # Construct specific target ID
            target_id = f"{resolved_law_id}_art_{art_num}_v_{resolved_version_name}"
            target_label = "ArticleVersion"
            
            if p_letter:
                target_id = f"{target_id}_p_{p_letter}"
                target_label = "Paragraph"
                if item_num:
                    target_id = f"{target_id}_i_{item_num}"
                    target_label = "Item"
                    
            citations.append({
                "ruling_id": extracted["ruling_id"],
                "target_id": target_id,
                "target_label": target_label,
                "citation_text": citation_context,
                "paragraph_letter": p_letter,
                "item_number": int(item_num) if item_num else None,
                "law_name": law_name.strip() if law_name else (resolved_law_id.replace("law_", "").replace("_", " "))
            })
            
        return {
            "judgment": extracted,
            "citations": citations
        }
